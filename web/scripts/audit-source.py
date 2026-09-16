"""Extract slide text and render selected PDF pages for migration audits.

This utility only reads the source deck/PDF and writes disposable evidence under
``web/tmp/source-audit``. It keeps the content migration reproducible without
modifying the original teaching material.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import pypdfium2 as pdfium
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_PPTX = ROOT / "docs" / "Modelos-fundacionales-en-vision-artificial.pptx"
DEFAULT_PDF = ROOT / "docs" / "Modelos-fundacionales-en-vision-artificial.pdf"
DEFAULT_OUT = ROOT / "web" / "tmp" / "source-audit"


def iter_shape_text(shape):
    if getattr(shape, "has_text_frame", False):
        text = "\n".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
        if text:
            yield text
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                yield " | ".join(cells)
    if getattr(shape, "shape_type", None) == 6:  # GROUP
        for child in shape.shapes:
            yield from iter_shape_text(child)


def slide_payload(prs: Presentation, number: int) -> dict:
    slide = prs.slides[number - 1]
    blocks = []
    for shape in sorted(slide.shapes, key=lambda item: (item.top, item.left)):
        blocks.extend(iter_shape_text(shape))
    notes = []
    if slide.has_notes_slide:
        for shape in slide.notes_slide.shapes:
            if getattr(shape, "has_text_frame", False):
                value = "\n".join(
                    p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()
                )
                if value and value not in {str(number), ""}:
                    notes.append(value)
    return {"slide": number, "text": blocks, "notes": notes}


def render_pages(pdf_path: Path, output: Path, start: int, end: int, scale: float = 1.5):
    pdf = pdfium.PdfDocument(str(pdf_path))
    rendered = []
    for number in range(start, end + 1):
        page = pdf[number - 1]
        bitmap = page.render(scale=scale)
        image = bitmap.to_pil().convert("RGB")
        target = output / f"slide-{number:02d}.png"
        image.save(target, optimize=True)
        rendered.append((number, target))
    return rendered


def contact_sheet(items, target: Path, columns: int = 3):
    if not items:
        return
    thumb_w, thumb_h, label_h = 480, 270, 34
    rows = (len(items) + columns - 1) // columns
    sheet = Image.new("RGB", (columns * thumb_w, rows * (thumb_h + label_h)), "white")
    draw = ImageDraw.Draw(sheet)
    font = ImageFont.load_default(size=18)
    for index, (number, path) in enumerate(items):
        image = Image.open(path).convert("RGB")
        image.thumbnail((thumb_w, thumb_h))
        x = (index % columns) * thumb_w + (thumb_w - image.width) // 2
        y = (index // columns) * (thumb_h + label_h)
        sheet.paste(image, (x, y))
        draw.text((x + 8, y + thumb_h + 6), f"Diapositiva {number}", fill="#16213a", font=font)
    sheet.save(target, optimize=True)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--start", type=int, default=31)
    parser.add_argument("--end", type=int, default=79)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()

    args.out.mkdir(parents=True, exist_ok=True)
    prs = Presentation(DEFAULT_PPTX)
    payload = [slide_payload(prs, n) for n in range(args.start, args.end + 1)]
    (args.out / "slides.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (args.out / "slides.txt").write_text(
        "\n\n".join(
            f"=== DIAPOSITIVA {item['slide']} ===\n" + "\n".join(item["text"])
            for item in payload
        ),
        encoding="utf-8",
    )
    rendered = render_pages(DEFAULT_PDF, args.out, args.start, args.end)
    ranges = {"tema-02": (31, 36), "tema-03": (37, 54), "tema-04": (55, 70), "tema-05": (71, 79)}
    for name, (start, end) in ranges.items():
        subset = [item for item in rendered if start <= item[0] <= end]
        contact_sheet(subset, args.out / f"{name}-contact.png")
    print(f"Audit evidence written to {args.out}")


if __name__ == "__main__":
    main()
